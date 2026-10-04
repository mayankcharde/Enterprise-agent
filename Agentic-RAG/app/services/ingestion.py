from pathlib import Path
from typing import Iterable
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from docx import Document as DocxDocument


SUPPORTED = {".pdf", ".txt", ".md", ".docx"}



def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        documents = PyPDFLoader(str(path)).load()
    elif suffix in {".txt", ".md"}:
        documents = TextLoader(str(path), encoding="utf-8").load()
    elif suffix == ".docx":
        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        documents = [Document(page_content=text, metadata={"source": str(path)})]
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    documents = [document for document in documents if document.page_content.strip()]
    for document in documents:
        document.metadata = {
            **document.metadata,
            "source": str(path),
            "source_type": "internal",
            "approved": True,
            "document_id": path.stem,
            "title": path.name,
        }
    if not documents:
        raise ValueError(f"No extractable text found in '{path}'. Add text or provide a text-based document.")
    return documents



def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    documents = [document for document in docs if document.page_content.strip()]
    if not documents:
        raise ValueError("No non-empty documents were provided for chunking.")
    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120, add_start_index=True)
    chunks = splitter.split_documents(documents)
    for index, chunk in enumerate(chunks):
        chunk.metadata = {
            **chunk.metadata,
            "chunk_id": str(index),
            "source_type": "internal",
            "approved": True,
        }
    if not chunks:
        raise ValueError("No chunks were created from the provided documents.")
    return chunks